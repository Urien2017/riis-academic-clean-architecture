from pathlib import Path
from tempfile import NamedTemporaryFile
from zipfile import ZIP_DEFLATED, ZipFile

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from xml.etree import ElementTree as ET


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "src" / "RIIS.Academic.Infrastructure" / "Documents" / "Templates" / "ProcesVerbalTemplate.docx"

BLUE = "1F4E79"
LIGHT_BLUE = "E8EEF5"
LIGHT_GRAY = "F4F6F8"
BORDER = "BFBFBF"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_borders(cell, color=BORDER, size="4"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.find(qn("w:tcBorders"))
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)

    for edge in ("top", "left", "bottom", "right"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_cell_width(cell, width_dxa):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width_dxa))
    tc_w.set(qn("w:type"), "dxa")


def set_table_width(table, width_dxa, column_widths):
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(width_dxa))
    tbl_w.set(qn("w:type"), "dxa")

    tbl_layout = tbl_pr.find(qn("w:tblLayout"))
    if tbl_layout is None:
        tbl_layout = OxmlElement("w:tblLayout")
        tbl_pr.append(tbl_layout)
    tbl_layout.set(qn("w:type"), "fixed")

    tbl_grid = table._tbl.tblGrid
    for child in list(tbl_grid):
        tbl_grid.remove(child)

    for width in column_widths:
        grid_col = OxmlElement("w:gridCol")
        grid_col.set(qn("w:w"), str(width))
        tbl_grid.append(grid_col)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    header = tr_pr.find(qn("w:tblHeader"))
    if header is None:
        header = OxmlElement("w:tblHeader")
        tr_pr.append(header)
    header.set(qn("w:val"), "true")


def normalize_cell_property_order(docx_path):
    ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
    order = {
        "cnfStyle": 0,
        "tcW": 1,
        "gridSpan": 2,
        "hMerge": 3,
        "vMerge": 4,
        "tcBorders": 5,
        "shd": 6,
        "noWrap": 7,
        "tcMar": 8,
        "textDirection": 9,
        "tcFitText": 10,
        "vAlign": 11,
        "hideMark": 12,
    }

    with ZipFile(docx_path, "r") as source, NamedTemporaryFile(delete=False, suffix=".docx") as tmp:
        tmp_path = Path(tmp.name)

        with ZipFile(tmp, "w", ZIP_DEFLATED) as target:
            for item in source.infolist():
                data = source.read(item.filename)

                if item.filename in {"word/document.xml", "word/header1.xml", "word/footer1.xml"}:
                    root = ET.fromstring(data)

                    for tc_pr in root.findall(".//w:tcPr", ns):
                        children = list(tc_pr)
                        children.sort(
                            key=lambda child: order.get(child.tag.rsplit("}", 1)[-1], 99)
                        )

                        for child in list(tc_pr):
                            tc_pr.remove(child)

                        for child in children:
                            tc_pr.append(child)

                    data = ET.tostring(root, encoding="utf-8", xml_declaration=True)

                target.writestr(item, data)

    docx_path.write_bytes(tmp_path.read_bytes())
    tmp_path.unlink(missing_ok=True)


def paragraph(text="", *, style=None, align=None, before=0, after=4, size=10, bold=False, color=None):
    p = doc.add_paragraph(style=style)
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.05
    if align is not None:
        p.alignment = align
    run = p.add_run(text)
    run.font.name = "Calibri"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    run.font.size = Pt(size)
    run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    return p


def fill_cell(cell, text, *, bold=False, color=None, fill=None, align=WD_ALIGN_PARAGRAPH.LEFT, size=9):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = align
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run(text)
    run.font.name = "Calibri"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    run.font.size = Pt(size)
    run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    if fill:
        set_cell_shading(cell, fill)
    set_cell_borders(cell)


doc = Document()
section = doc.sections[0]
section.orientation = WD_ORIENT.LANDSCAPE
section.page_width = Cm(29.7)
section.page_height = Cm(21.0)
section.top_margin = Cm(1.0)
section.bottom_margin = Cm(1.0)
section.left_margin = Cm(0.8)
section.right_margin = Cm(0.8)
section.header_distance = Cm(0.5)
section.footer_distance = Cm(0.5)

style = doc.styles["Normal"]
style.font.name = "Calibri"
style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
style.font.size = Pt(10)

header = section.header.paragraphs[0]
header.text = "RIIS Academic - Procès-verbal des notes"
header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
header.runs[0].font.size = Pt(8)
header.runs[0].font.color.rgb = RGBColor(90, 90, 90)

footer = section.footer.paragraphs[0]
footer.text = "Document généré automatiquement par RIIS Academic"
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
footer.runs[0].font.size = Pt(8)
footer.runs[0].font.color.rgb = RGBColor(90, 90, 90)

paragraph("PROCES VERBAL DES NOTES", align=WD_ALIGN_PARAGRAPH.CENTER, size=16, bold=True, color=BLUE, after=2)
paragraph("{{PV_TYPE}} - {{SEMESTRE_LIBELLE}}", align=WD_ALIGN_PARAGRAPH.CENTER, size=11, bold=True, color=BLUE, after=8)

meta_widths = [1800, 5200, 1800, 5200]
meta = doc.add_table(rows=4, cols=4)
meta.alignment = WD_TABLE_ALIGNMENT.CENTER
set_table_width(meta, sum(meta_widths), meta_widths)
labels_values = [
    ("Année académique", "{{ANNEE_ACADEMIQUE}}", "Cycle", "{{CYCLE_FORMATION}}"),
    ("Filière / Spécialité", "{{PARCOURS}}", "Niveau", "{{NIVEAU}}"),
    ("Classe", "{{CLASSE_PEDAGOGIQUE}}", "Semestre", "{{SEMESTRE_NUMERO}}"),
    ("Session", "{{SESSION}}", "Statut", "{{STATUT}}"),
]
for row, values in zip(meta.rows, labels_values):
    for idx, value in enumerate(values):
        is_label = idx in (0, 2)
        fill_cell(
            row.cells[idx],
            value,
            bold=is_label,
            fill=LIGHT_GRAY if is_label else None,
            size=8.5,
        )
        set_cell_width(row.cells[idx], meta_widths[idx])

paragraph("", after=2)
paragraph(
    "Le tableau détaillé du PV sera inséré automatiquement à l’emplacement ci-dessous.",
    align=WD_ALIGN_PARAGRAPH.LEFT,
    size=8,
    color="666666",
    after=2,
)

placeholder_table = doc.add_table(rows=2, cols=1)
placeholder_table.alignment = WD_TABLE_ALIGNMENT.CENTER
set_table_width(placeholder_table, 15600, [15600])
fill_cell(placeholder_table.rows[0].cells[0], "{{PV_TABLE}}", bold=True, color=BLUE, fill=LIGHT_BLUE, align=WD_ALIGN_PARAGRAPH.CENTER, size=10)
fill_cell(
    placeholder_table.rows[1].cells[0],
    "Placeholder technique : ne pas supprimer. Le service d’export remplacera cette zone par le tableau des étudiants, EC, CC/CCON/SN, moyennes, crédits, rangs et décisions.",
    color="666666",
    align=WD_ALIGN_PARAGRAPH.CENTER,
    size=8,
)

paragraph("", after=4)

summary_widths = [2500, 1800, 2500, 1800, 2500, 1800]
summary = doc.add_table(rows=1, cols=6)
summary.alignment = WD_TABLE_ALIGNMENT.CENTER
set_table_width(summary, sum(summary_widths), summary_widths)
summary_values = [
    ("Étudiants", "{{NOMBRE_ETUDIANTS}}"),
    ("Validés", "{{NOMBRE_VALIDES}}"),
    ("Rattrapage", "{{NOMBRE_RATTRAPAGE}}"),
]
index = 0
for label, value in summary_values:
    fill_cell(summary.rows[0].cells[index], label, bold=True, fill=LIGHT_GRAY, size=8.5)
    set_cell_width(summary.rows[0].cells[index], summary_widths[index])
    index += 1
    fill_cell(summary.rows[0].cells[index], value, align=WD_ALIGN_PARAGRAPH.CENTER, size=8.5)
    set_cell_width(summary.rows[0].cells[index], summary_widths[index])
    index += 1

paragraph("", after=8)
paragraph("Observations : {{OBSERVATION}}", size=9, after=12)

signature_widths = [5200, 5200]
sign = doc.add_table(rows=2, cols=2)
sign.alignment = WD_TABLE_ALIGNMENT.CENTER
set_table_width(sign, sum(signature_widths), signature_widths)
for idx, title in enumerate(("Le responsable académique", "Le jury")):
    fill_cell(sign.rows[0].cells[idx], title, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, size=9)
    set_cell_width(sign.rows[0].cells[idx], signature_widths[idx])
    fill_cell(sign.rows[1].cells[idx], "\n\n____________________________", align=WD_ALIGN_PARAGRAPH.CENTER, size=9)
    set_cell_width(sign.rows[1].cells[idx], signature_widths[idx])

OUTPUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUTPUT)
normalize_cell_property_order(OUTPUT)
print(OUTPUT)

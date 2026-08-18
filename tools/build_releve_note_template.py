from pathlib import Path
from tempfile import NamedTemporaryFile
from zipfile import ZIP_DEFLATED, ZipFile

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from xml.etree import ElementTree as ET


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "src" / "RIIS.Academic.Infrastructure" / "Documents" / "Templates" / "ReleveNoteTemplate.docx"

BLUE = "1F4E79"
LIGHT_BLUE = "E8EEF5"
LIGHT_GRAY = "F4F6F8"
BORDER = "BFBFBF"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_borders(cell, color=BORDER, size="4"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.find(qn("w:tcBorders"))
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)

    for edge in ("top", "left", "bottom", "right"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_cell_width(cell, width_dxa):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width_dxa))
    tc_w.set(qn("w:type"), "dxa")


def set_table_width(table, width_dxa, column_widths):
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(width_dxa))
    tbl_w.set(qn("w:type"), "dxa")

    tbl_jc = tbl_pr.find(qn("w:jc"))
    if tbl_jc is None:
        tbl_jc = OxmlElement("w:jc")
        tbl_pr.append(tbl_jc)
    tbl_jc.set(qn("w:val"), "center")

    tbl_layout = tbl_pr.find(qn("w:tblLayout"))
    if tbl_layout is None:
        tbl_layout = OxmlElement("w:tblLayout")
        tbl_pr.append(tbl_layout)
    tbl_layout.set(qn("w:type"), "fixed")

    tbl_grid = table._tbl.tblGrid
    for child in list(tbl_grid):
        tbl_grid.remove(child)

    for width in column_widths:
        grid_col = OxmlElement("w:gridCol")
        grid_col.set(qn("w:w"), str(width))
        tbl_grid.append(grid_col)


def normalize_cell_property_order(docx_path):
    ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
    order = {
        "cnfStyle": 0,
        "tcW": 1,
        "gridSpan": 2,
        "hMerge": 3,
        "vMerge": 4,
        "tcBorders": 5,
        "shd": 6,
        "noWrap": 7,
        "tcMar": 8,
        "textDirection": 9,
        "tcFitText": 10,
        "vAlign": 11,
        "hideMark": 12,
    }

    with ZipFile(docx_path, "r") as source, NamedTemporaryFile(delete=False, suffix=".docx") as tmp:
        tmp_path = Path(tmp.name)

        with ZipFile(tmp, "w", ZIP_DEFLATED) as target:
            for item in source.infolist():
                data = source.read(item.filename)

                if item.filename in {"word/document.xml", "word/header1.xml", "word/footer1.xml"}:
                    root = ET.fromstring(data)

                    for tc_pr in root.findall(".//w:tcPr", ns):
                        children = list(tc_pr)
                        children.sort(
                            key=lambda child: order.get(child.tag.rsplit("}", 1)[-1], 99)
                        )

                        for child in list(tc_pr):
                            tc_pr.remove(child)

                        for child in children:
                            tc_pr.append(child)

                    data = ET.tostring(root, encoding="utf-8", xml_declaration=True)

                target.writestr(item, data)

    docx_path.write_bytes(tmp_path.read_bytes())
    tmp_path.unlink(missing_ok=True)


def paragraph(text="", *, align=None, before=0, after=4, size=10, bold=False, color=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.05
    p.alignment = align if align is not None else WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text)
    run.font.name = "Calibri"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    run.font.size = Pt(size)
    run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    return p


def fill_cell(cell, text, *, bold=False, color=None, fill=None, align=WD_ALIGN_PARAGRAPH.LEFT, size=9):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = align
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run(text)
    run.font.name = "Calibri"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    run.font.size = Pt(size)
    run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    if fill:
        set_cell_shading(cell, fill)
    set_cell_borders(cell)


doc = Document()
section = doc.sections[0]
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.top_margin = Inches(0.5)
section.bottom_margin = Inches(0.5)
section.left_margin = Inches(0.5)
section.right_margin = Inches(0.5)
section.header_distance = Inches(0.3)
section.footer_distance = Inches(0.3)

style = doc.styles["Normal"]
style.font.name = "Calibri"
style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
style.font.size = Pt(10)

header = section.header.paragraphs[0]
header.text = "RIIS Academic - Relevé de notes"
header.alignment = WD_ALIGN_PARAGRAPH.CENTER
header.runs[0].font.size = Pt(8)
header.runs[0].font.color.rgb = RGBColor(90, 90, 90)

footer = section.footer.paragraphs[0]
footer.text = "Document généré automatiquement par RIIS Academic"
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
footer.runs[0].font.size = Pt(8)
footer.runs[0].font.color.rgb = RGBColor(90, 90, 90)

paragraph("RÉPUBLIQUE DU CAMEROUN", align=WD_ALIGN_PARAGRAPH.CENTER, size=9, after=0)
paragraph("Paix - Travail - Patrie", align=WD_ALIGN_PARAGRAPH.CENTER, size=9, after=4)
paragraph("{{TITRE}}", align=WD_ALIGN_PARAGRAPH.CENTER, size=15, bold=True, color=BLUE, after=1)
paragraph("{{CYCLE_FORMATION}} - {{ANNEE_ACADEMIQUE}}", align=WD_ALIGN_PARAGRAPH.CENTER, size=11, bold=True, color=BLUE, after=8)

info_widths = [2100, 7260]
info = doc.add_table(rows=8, cols=2)
info.alignment = WD_TABLE_ALIGNMENT.CENTER
set_table_width(info, sum(info_widths), info_widths)

labels_values = [
    ("Nom et prénoms", "{{ETUDIANT_NOM_COMPLET}}"),
    ("Matricule", "{{MATRICULE}}"),
    ("Né(e) le", "{{DATE_NAISSANCE}}"),
    ("Lieu de naissance", "{{LIEU_NAISSANCE}}"),
    ("Filière", "{{FILIERE}}"),
    ("Spécialité", "{{SPECIALITE}}"),
    ("Niveau", "{{NIVEAU}}"),
    ("Année académique", "{{ANNEE_ACADEMIQUE}}"),
]

for row, (label, value) in zip(info.rows, labels_values):
    fill_cell(row.cells[0], label, bold=True, fill=LIGHT_GRAY, size=8.5)
    fill_cell(row.cells[1], value, size=8.5)
    set_cell_width(row.cells[0], info_widths[0])
    set_cell_width(row.cells[1], info_widths[1])

paragraph("", after=2)
paragraph(
    "Les tableaux des semestres, UE, EC, crédits et décisions seront insérés automatiquement ci-dessous.",
    align=WD_ALIGN_PARAGRAPH.CENTER,
    size=8,
    color="666666",
    after=2,
)

placeholder_table = doc.add_table(rows=2, cols=1)
placeholder_table.alignment = WD_TABLE_ALIGNMENT.CENTER
set_table_width(placeholder_table, 9360, [9360])
fill_cell(
    placeholder_table.rows[0].cells[0],
    "{{RELEVE_TABLES}}",
    bold=True,
    color=BLUE,
    fill=LIGHT_BLUE,
    align=WD_ALIGN_PARAGRAPH.CENTER,
    size=10,
)
fill_cell(
    placeholder_table.rows[1].cells[0],
    "Placeholder technique : ne pas supprimer. Le service d’export remplacera cette zone par les tableaux du relevé.",
    color="666666",
    align=WD_ALIGN_PARAGRAPH.CENTER,
    size=8,
)

paragraph("", after=4)

decision_widths = [2200, 2480, 2200, 2480]
decision = doc.add_table(rows=3, cols=4)
decision.alignment = WD_TABLE_ALIGNMENT.CENTER
set_table_width(decision, sum(decision_widths), decision_widths)
decision_values = [
    ("Moyenne annuelle", "{{MOYENNE_ANNUELLE}}", "Mention", "{{MENTION}}"),
    ("Crédits capitalisés", "{{CREDITS_CAPITALISES}}", "Crédits requis", "{{CREDITS_REQUIS}}"),
    ("Décision", "{{DECISION}}", "Date d’édition", "{{DATE_EDITION}}"),
]

for row, values in zip(decision.rows, decision_values):
    for idx, value in enumerate(values):
        is_label = idx in (0, 2)
        fill_cell(
            row.cells[idx],
            value,
            bold=is_label,
            fill=LIGHT_GRAY if is_label else None,
            align=WD_ALIGN_PARAGRAPH.LEFT if is_label else WD_ALIGN_PARAGRAPH.CENTER,
            size=8.5,
        )
        set_cell_width(row.cells[idx], decision_widths[idx])

paragraph("", after=8)
paragraph("Fait à Yaoundé, le {{DATE_EDITION}}", align=WD_ALIGN_PARAGRAPH.CENTER, size=9, after=16)
paragraph("Le Directeur", align=WD_ALIGN_PARAGRAPH.CENTER, size=9, bold=True, after=28)
paragraph("____________________________", align=WD_ALIGN_PARAGRAPH.CENTER, size=9, after=0)

OUTPUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUTPUT)
normalize_cell_property_order(OUTPUT)
print(OUTPUT)

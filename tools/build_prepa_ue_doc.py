# -*- coding: utf-8 -*-
from __future__ import annotations

import os
import re
import unicodedata
import zipfile
from pathlib import Path
from xml.etree import ElementTree as ET

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

NS = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}


def normalize(value: str | None) -> str:
    value = (value or "").lower().strip()
    value = "".join(
        char
        for char in unicodedata.normalize("NFD", value)
        if unicodedata.category(char) != "Mn"
    )
    value = re.sub(r"[^a-z0-9]+", " ", value)
    return re.sub(r"\s+", " ", value).strip()


UE_BY_COURSE = {
    # Management / Economie
    "climate and environment economic news": "Environnement économique et actualité",
    "business computer science 1": "Informatique et outils numériques",
    "reglementation": "Droit, réglementation et éthique",
    "mathematiques pour l economie et la gestion": "Mathématiques et techniques quantitatives",
    "francais": "Formation bilingue et communication",
    "marketing fondamentale": "Management, marketing et entrepreneuriat",
    "comptabilites generale general accounting": "Comptabilité, finance et gestion",
    "applied ethics": "Droit, réglementation et éthique",
    "microeconomie": "Economie générale",
    "communication interculturelle": "Formation bilingue et communication",
    "introduction au droit": "Droit, réglementation et éthique",
    "introduction au management": "Management, marketing et entrepreneuriat",
    "intermediate english": "Formation bilingue et communication",
    "projet professionnel et personnel": "Projet professionnel et personnel",
    "comptabilites financiere": "Comptabilité, finance et gestion",
    "business computer science 2": "Informatique et outils numériques",
    "administration": "Management, marketing et entrepreneuriat",
    "outils d analyse quantitative de gestion": "Mathématiques et techniques quantitatives",
    "competence linguistique": "Formation bilingue et communication",
    "initiation au marketing digitale": "Management, marketing et entrepreneuriat",
    "comptabilites 2": "Comptabilité, finance et gestion",
    "statistiques descriptive et probabilite": "Mathématiques et techniques quantitatives",
    "macroeconomie": "Economie générale",
    "identite numerique competences digitales": "Informatique et outils numériques",
    "element de droit des affaires droit des societes commerciales": "Droit, réglementation et éthique",
    "management des organisations": "Management, marketing et entrepreneuriat",
    "concepts et actualites economiques": "Environnement économique et actualité",
    "business english": "Formation bilingue et communication",
    "comptabilite des societes": "Comptabilité, finance et gestion",
    "mathematiques financieres": "Mathématiques et techniques quantitatives",
    "informatiques de gestion 2": "Informatique et outils numériques",
    "business communication": "Formation bilingue et communication",
    "e marketing": "Management, marketing et entrepreneuriat",
    "international business": "Environnement économique et actualité",
    "business law": "Droit, réglementation et éthique",
    "financial analysis": "Comptabilité, finance et gestion",
    "comptabilite analytique": "Comptabilité, finance et gestion",
    "projet personnel et professionnel": "Projet professionnel et personnel",
    "global ethics for contemporary societies": "Droit, réglementation et éthique",
    "entrepreneurship development": "Management, marketing et entrepreneuriat",
    "controle de gestion": "Comptabilité, finance et gestion",
    "management des projets": "Management, marketing et entrepreneuriat",
    "finance des entreprises diagnostique financiers": "Comptabilité, finance et gestion",
    "rse gestion previsionnelle des entreprises": "Management, marketing et entrepreneuriat",
    "systeme d information": "Informatique et outils numériques",
    "developpementpersonnel relations professionnelle": "Projet professionnel et personnel",
    "logique et raisonnement": "Mathématiques et techniques quantitatives",
    "fiscalite": "Droit, réglementation et éthique",
    "relations internationales": "Environnement économique et actualité",
    "c a o gestion de la paie economie d entreprise": "Comptabilité, finance et gestion",
    "stage": "Stage professionnel",
    # Ingénieur
    "electricite et electronique": "Electricité, électronique et systèmes",
    "mecanique": "Sciences fondamentales pour l’ingénieur",
    "numerical data bases and application": "Informatique, données et programmation",
    "algorithme et application": "Informatique, données et programmation",
    "mathematiques appliquees pour l ingenieur": "Mathématiques appliquées",
    "basics mathematics": "Mathématiques appliquées",
    "chimie generale": "Sciences physiques et chimiques",
    "communication": "Formation bilingue et communication",
    "physique des ondes et traitement du signal": "Sciences physiques et traitement du signal",
    "travaux pratique informatique": "Informatique, données et programmation",
    "physique 2": "Sciences physiques et chimiques",
    "outils et methodes de l ingenieur": "Méthodes et processus industriels",
    "processus industriels": "Méthodes et processus industriels",
    "probabilites et statistiques": "Mathématiques appliquées",
    "reglementation 2": "Droit, réglementation et éthique",
    "travaux pratique de physiques": "Sciences physiques et chimiques",
    "bases de donnees": "Informatique, données et programmation",
    "circuit logique et algebre de boole": "Electricité, électronique et systèmes",
    "mathematiques appliquees pour l ingenieur 2": "Mathématiques appliquées",
    "creation et organisation des entreprises": "Management, entrepreneuriat et organisation",
    "mecanique des fluides": "Sciences fondamentales pour l’ingénieur",
    "base en programmation programmation orientee objet": "Informatique, données et programmation",
    "architecture du reseau local": "Réseaux, systèmes et architecture informatique",
    "base de donnees": "Informatique, données et programmation",
    "architecture des systemes informatique": "Réseaux, systèmes et architecture informatique",
    "initiation au genie logiciel": "Informatique, données et programmation",
    "anglais": "Formation bilingue et communication",
    "initiation au projet": "Projet professionnel et personnel",
    "electrotechnique embarque": "Electricité, électronique et systèmes",
    "algorithme et langage": "Informatique, données et programmation",
    "architecture distribuees geotechnique": "Architecture, construction et géotechnique",
    "programmation web rdm": "Informatique, programmation et résistance des matériaux",
    "systeme linux reseau et administration systeme et fondation et soutenement": "Réseaux, systèmes et cybersécurité",
    "cybersecurite topographie et essaie": "Réseaux, systèmes et cybersécurité",
    "algebre multitechnique": "Mathématiques appliquées",
    "initiation a la fabrication multitechnique": "Méthodes et processus industriels",
    "projet dao cao bureau d etude et ouvrage d art": "Conception, DAO/CAO et bureau d’études",
    "developpement personnel relation professionnelle": "Projet professionnel et personnel",
    "terminaux mobile uml beton arme": "Technologies mobiles, UML et structures",
    "ethique des technologies et du numerique": "Droit, réglementation et éthique",
}


def extract_prepa_rows(source_path: Path) -> list[dict[str, str]]:
    with zipfile.ZipFile(source_path) as archive:
        root = ET.fromstring(archive.read("word/document.xml"))

    table = root.findall(".//w:tbl", NS)[0]
    rows: list[list[str]] = []
    for table_row in table.findall("./w:tr", NS):
        cells = []
        for table_cell in table_row.findall("./w:tc", NS):
            parts = []
            for paragraph in table_cell.findall(".//w:p", NS):
                text = "".join(
                    text_node.text or ""
                    for text_node in paragraph.findall(".//w:t", NS)
                ).strip()
                if text:
                    parts.append(text)
            cells.append(" / ".join(parts).strip())
        rows.append(cells)

    data: list[dict[str, str]] = []
    specialite = ""
    semestre = ""
    for cells in rows:
        joined = " ".join(cell.strip() for cell in cells).strip()
        normalized_joined = normalize(joined)
        if not joined:
            continue

        if len(cells) == 1 and all(
            marker not in normalized_joined
            for marker in ["total", "semestre", "cours"]
        ):
            specialite = joined
            continue

        if cells and normalize(cells[0]).startswith("semestre"):
            semestre = cells[0].strip()

        cours = cells[1].strip() if len(cells) > 1 else ""
        heures = cells[2].strip() if len(cells) > 2 else ""
        credits = cells[3].strip() if len(cells) > 3 else ""

        if cours and "cours" not in normalize(cours) and "total" not in normalize(cours):
            data.append(
                {
                    "specialite": specialite,
                    "semestre": semestre,
                    "cours": cours,
                    "heures": heures,
                    "credits": credits,
                }
            )

    return data


def resolve_ue(row: dict[str, str]) -> str:
    course_key = normalize(row["cours"])
    if course_key == "global ethics for contemporary societies":
        return "Droit, réglementation et éthique"
    if course_key == "reglementation":
        return "Droit, réglementation et éthique"

    ue = UE_BY_COURSE.get(course_key)
    if not ue:
        raise KeyError(f"Aucune UE pour le cours: {row['cours']} ({course_key})")
    return ue


def set_cell_text(cell, text: str, *, bold=False, size=8.0, color=None, align=None):
    cell.text = ""
    paragraph = cell.paragraphs[0]
    if align is not None:
        paragraph.alignment = align
    run = paragraph.add_run(text or "")
    run.bold = bold
    run.font.name = "Calibri"
    run.font.size = Pt(size)
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def shade_cell(cell, fill: str):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, *, top=80, start=80, bottom=80, end=80):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)

    for margin, value in {
        "top": top,
        "start": start,
        "bottom": bottom,
        "end": end,
    }.items():
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    row_props = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    row_props.append(header)


def set_table_grid(table, widths: list[int]):
    table_xml = table._tbl
    table_props = table_xml.tblPr

    table_width = table_props.find(qn("w:tblW"))
    if table_width is None:
        table_width = OxmlElement("w:tblW")
        table_props.append(table_width)
    table_width.set(qn("w:w"), str(sum(widths)))
    table_width.set(qn("w:type"), "dxa")

    table_indent = table_props.find(qn("w:tblInd"))
    if table_indent is None:
        table_indent = OxmlElement("w:tblInd")
        table_props.append(table_indent)
    table_indent.set(qn("w:w"), "120")
    table_indent.set(qn("w:type"), "dxa")

    layout = table_props.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        table_props.append(layout)
    layout.set(qn("w:type"), "fixed")

    old_grid = table_xml.find(qn("w:tblGrid"))
    if old_grid is not None:
        table_xml.remove(old_grid)

    grid = OxmlElement("w:tblGrid")
    for width in widths:
        column = OxmlElement("w:gridCol")
        column.set(qn("w:w"), str(width))
        grid.append(column)
    table_xml.insert(0, grid)

    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            tc_pr = cell._tc.get_or_add_tcPr()
            cell_width = tc_pr.find(qn("w:tcW"))
            if cell_width is None:
                cell_width = OxmlElement("w:tcW")
                tc_pr.append(cell_width)
            cell_width.set(qn("w:w"), str(width))
            cell_width.set(qn("w:type"), "dxa")


def merge_repeated_cells(table, detail_rows: list[dict]):
    """Merge consecutive repeated values in Spécialité, Semestre and UE columns.

    The source order is preserved. Only adjacent repeated values inside the same
    spécialité/semestre block are merged, which keeps the document faithful to
    the original programme while removing visual repetition.
    """

    for column_index in [0, 1, 2]:
        start = 0
        while start < len(detail_rows):
            current = detail_rows[start]
            value = current["values"][column_index]
            group = current["group"]
            end = start

            while (
                end + 1 < len(detail_rows)
                and detail_rows[end + 1]["group"] == group
                and detail_rows[end + 1]["values"][column_index] == value
            ):
                end += 1

            if end > start:
                first_cell = table.cell(current["row_index"], column_index)
                last_cell = table.cell(detail_rows[end]["row_index"], column_index)
                merged_cell = first_cell.merge(last_cell)
                merged_cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                set_cell_margins(merged_cell, top=70, bottom=70, start=90, end=90)
                align = (
                    WD_ALIGN_PARAGRAPH.CENTER
                    if column_index in [0, 1]
                    else WD_ALIGN_PARAGRAPH.LEFT
                )
                set_cell_text(merged_cell, value, size=7.5, align=align)

            start = end + 1


def build_document(source_path: Path, output_path: Path):
    data = extract_prepa_rows(source_path)
    # Validate every EC before writing the file.
    for row in data:
        resolve_ue(row)

    document = Document()
    section = document.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width = Inches(11)
    section.page_height = Inches(8.5)
    section.top_margin = Inches(0.55)
    section.bottom_margin = Inches(0.55)
    section.left_margin = Inches(0.55)
    section.right_margin = Inches(0.55)

    document.styles["Normal"].font.name = "Calibri"
    document.styles["Normal"].font.size = Pt(9)

    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title.add_run("Programme des spécialités en PREPA avec Unités d’enseignement (UE)")
    title_run.bold = True
    title_run.font.name = "Calibri"
    title_run.font.size = Pt(16)
    title_run.font.color.rgb = RGBColor(31, 77, 120)

    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle_run = subtitle.add_run(
        "Document dérivé du programme source PREPA - colonne UE ajoutée, intitulés EC conservés exactement"
    )
    subtitle_run.italic = True
    subtitle_run.font.name = "Calibri"
    subtitle_run.font.size = Pt(9)
    subtitle_run.font.color.rgb = RGBColor(89, 89, 89)

    headers = ["Spécialité", "Semestre", "UE", "EC / Cours", "Nbre d’heures", "Crédits"]
    table = document.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"

    set_repeat_table_header(table.rows[0])
    for index, header in enumerate(headers):
        cell = table.rows[0].cells[index]
        shade_cell(cell, "E8EEF5")
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_margins(cell, top=80, bottom=80, start=100, end=100)
        set_cell_text(
            cell,
            header,
            bold=True,
            size=8,
            color="1F4D78",
            align=WD_ALIGN_PARAGRAPH.CENTER,
        )

    current_group = None
    detail_rows = []
    for row_data in data:
        group = (row_data["specialite"], row_data["semestre"])
        if group != current_group:
            group_row = table.add_row()
            merged = group_row.cells[0].merge(group_row.cells[-1])
            shade_cell(merged, "F4F6F9")
            set_cell_margins(merged, top=90, bottom=90, start=120, end=120)
            set_cell_text(
                merged,
                f"{row_data['specialite']} - {row_data['semestre']}",
                bold=True,
                size=8.5,
                color="1F4D78",
                align=WD_ALIGN_PARAGRAPH.LEFT,
            )
            current_group = group

        table_row = table.add_row()
        semestre = row_data["semestre"].replace("SEMESTRE ", "S")
        values = [
            row_data["specialite"],
            semestre,
            resolve_ue(row_data),
            row_data["cours"],
            row_data["heures"],
            row_data["credits"],
        ]
        for cell, value in zip(table_row.cells, values):
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell, top=70, bottom=70, start=90, end=90)
            align = (
                WD_ALIGN_PARAGRAPH.CENTER
                if value in [row_data["specialite"], semestre, row_data["heures"], row_data["credits"]]
                else WD_ALIGN_PARAGRAPH.LEFT
            )
            set_cell_text(cell, value, size=7.5, align=align)

        detail_rows.append(
            {
                "group": group,
                "row_index": len(table.rows) - 1,
                "values": values,
            }
        )

    merge_repeated_cells(table, detail_rows)
    set_table_grid(table, [1450, 720, 2400, 4300, 820, 650])

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer_run = footer.add_run("Programme PREPA avec UE - document de travail pour validation avant seeder officiel")
    footer_run.font.name = "Calibri"
    footer_run.font.size = Pt(8)
    footer_run.font.color.rgb = RGBColor(89, 89, 89)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    document.save(output_path)
    return data


if __name__ == "__main__":
    source = Path(os.environ["PREPA_DOCX"])
    output = Path(os.environ["OUT_DOCX"])
    rows = build_document(source, output)
    print(output)
    print(f"lignes_ec={len(rows)}")
    print(f"ues={len({resolve_ue(row) for row in rows})}")
